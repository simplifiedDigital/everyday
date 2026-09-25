import contextlib
import csv
import hashlib
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import wave
import zipfile

from PIL import Image
import pymupdf
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from engine import worker


class ToolsTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.out = self.root / 'outputs'
        self.pdf = self.root / 'Example.pdf'
        with pymupdf.open() as doc:
            for n in range(3):
                page = doc.new_page(width=400, height=500)
                page.insert_text((40, 50), f'Page {n+1}', fontsize=20)
                page.insert_text((40, 90), 'PRIVATE 123456', fontsize=12)
                page.insert_text((40, 160), 'Keep this paragraph.', fontsize=12)
            doc.save(self.pdf)
        self.photo = self.root / 'Photo.png'
        Image.new('RGBA', (300, 200), (200, 30, 50, 120)).save(self.photo)
        self.original_hash = worker.hash_file(self.pdf)
        self.protocol = worker.PROTOCOL; worker.PROTOCOL = io.StringIO()

    def tearDown(self):
        worker.PROTOCOL = self.protocol
        worker.REQUEST = {}
        self.assertEqual(self.original_hash, worker.hash_file(self.pdf))
        self.temp.cleanup()

    def run_tool(self, tool, files=None, **options):
        return worker.dispatch({'tool': tool, 'files': [str(p) for p in (files if files is not None else [self.pdf])], 'options': options, 'output_dir': str(self.out)})

    def test_split_merge_extract(self):
        split = self.run_tool('pdf-split', pages='1,3')
        self.assertEqual(len(split['files']), 2)
        merged = self.run_tool('pdf-merge', files=split['files'])
        with pymupdf.open(merged['files'][0]) as doc:
            self.assertEqual(len(doc), 2)
            self.assertIn('Page 3', doc[1].get_text())
        extract = self.run_tool('pdf-extract', page=2)
        with pymupdf.open(extract['files'][0]) as doc:
            self.assertEqual(len(doc), 1); self.assertIn('Page 2', doc[0].get_text())

    def test_page_range_validation(self):
        for bad in ['0', '4', '3-1', 'one', '1-2-3']:
            with self.assertRaises(worker.UserError): self.run_tool('pdf-split', pages=bad)

    def test_rotation(self):
        result = self.run_tool('pdf-rotate', pages='2', angle=90)
        with pymupdf.open(result['files'][0]) as doc:
            self.assertEqual([p.rotation for p in doc], [0,90,0])

    def test_password_roundtrip(self):
        locked = self.run_tool('pdf-protect', newPassword='a-real-test-password')['files'][0]
        with pymupdf.open(locked) as doc: self.assertTrue(doc.needs_pass)
        with self.assertRaises(worker.UserError): self.run_tool('pdf-unlock', files=[locked], password='wrong')
        unlocked = self.run_tool('pdf-unlock', files=[locked], password='a-real-test-password')['files'][0]
        with pymupdf.open(unlocked) as doc: self.assertFalse(doc.needs_pass); self.assertIn('PRIVATE', doc[0].get_text())

    def test_redaction_removes_underlying_text(self):
        result = self.run_tool('pdf-redact', boxes=[{'page': 0, 'x': .08, 'y': .14, 'w': .65, 'h': .08}])
        with pymupdf.open(result['files'][0]) as doc:
            self.assertNotIn('PRIVATE', doc[0].get_text())
            self.assertIn('Keep this paragraph', doc[0].get_text())
            self.assertIn('PRIVATE', doc[1].get_text())
            for xref in doc[0].get_contents():
                stream = doc.xref_stream(xref)
                if stream: self.assertNotIn(b'PRIVATE', stream)

    def test_redaction_rotated_coordinates(self):
        rotated = self.run_tool('pdf-rotate', pages='1', angle=90)['files'][0]
        with pymupdf.open(rotated) as doc:
            p = doc[0]; r = p.search_for('PRIVATE 123456')[0] * p.rotation_matrix
            box = {'page': 0, 'x': r.x0/p.rect.width, 'y': r.y0/p.rect.height, 'w': r.width/p.rect.width, 'h': r.height/p.rect.height}
        output = self.run_tool('pdf-redact', files=[rotated], boxes=[box])['files'][0]
        with pymupdf.open(output) as doc: self.assertNotIn('PRIVATE', doc[0].get_text())

    def test_markdown_with_images_and_scan(self):
        source = self.root / 'illustrated.pdf'
        with pymupdf.open() as doc:
            p = doc.new_page(); p.insert_text((50,50), 'My heading', fontsize=24)
            p.insert_text((50,95), 'A paragraph with an image below.', fontsize=12)
            p.insert_image(pymupdf.Rect(50,120,350,320), filename=self.photo)
            p = doc.new_page(); p.insert_image(pymupdf.Rect(0,0,300,200), filename=self.photo)
            doc.save(source)
        result = self.run_tool('pdf-markdown', files=[source])
        md = Path(result['files'][0]).read_text(encoding='utf-8')
        self.assertIn('My heading', md); self.assertIn('images/', md)
        self.assertNotIn(str(self.root), md)
        self.assertTrue(list((Path(result['folder']) / 'images').glob('*.png')))
        self.assertIn('scanned', result['warning'])

    def test_resize_conversion_and_collage(self):
        resized = self.run_tool('image-resize', files=[self.photo], width=120, format='jpg')
        with Image.open(resized['files'][0]) as img: self.assertEqual(img.size, (120,80)); self.assertEqual(img.mode, 'RGB')
        converted = self.run_tool('image-convert', files=[self.photo], format='webp')
        with Image.open(converted['files'][0]) as img: self.assertEqual(img.format, 'WEBP')
        collage = self.run_tool('image-collage', files=[self.photo, self.photo], columns=2)
        with Image.open(collage['files'][0]) as img: self.assertEqual(img.size, (1328,672))
        preview = self.run_tool('inspect', files=[self.photo, self.photo], collage=True, columns=2)
        self.assertEqual((preview['width'], preview['height']), (332,168))

    def test_zip_roundtrip(self):
        archive = self.run_tool('file-zip', files=[self.photo, self.pdf])['files'][0]
        output = self.run_tool('file-unzip', files=[archive])
        restored = Path(output['folder']) / self.pdf.name
        self.assertEqual(worker.hash_file(restored), self.original_hash)

    def test_zip_traversal_rejected(self):
        source = self.root / 'unsafe.zip'
        with zipfile.ZipFile(source, 'w') as z: z.writestr('../escape.txt', 'bad')
        with self.assertRaises(worker.UserError): self.run_tool('file-unzip', files=[source])
        self.assertFalse((self.root / 'escape.txt').exists())

    def test_split_join_and_corrupt_part(self):
        source = self.root / 'large.bin'; source.write_bytes(b'0123456789' * 250000)
        split = self.run_tool('file-split', files=[source], size=1)
        joined = self.run_tool('file-join', files=[split['files'][0]])
        self.assertEqual(worker.hash_file(source), worker.hash_file(joined['files'][0]))
        Path(split['files'][1]).write_bytes(b'corrupt')
        with self.assertRaises(worker.UserError): self.run_tool('file-join', files=[split['files'][0]])

    def make_excel(self):
        path = self.root / 'Book.xlsx'; wb = Workbook(); ws = wb.active; ws.title = 'Sales'
        ws.append(['Name', 'Amount']); ws.append(['Alice', 12]); ws['B2'].number_format = '$#,##0.00'
        ws['A1'].font = Font(bold=True, color='123456'); wb.save(path); wb.close(); return path

    def test_excel_csv_values(self):
        result = self.run_tool('sheet-csv', files=[self.make_excel()])
        with open(result['files'][0], encoding='utf-8-sig') as f: rows = list(csv.reader(f))
        self.assertEqual(rows, [['Name','Amount'],['Alice','12']])

    def test_merge_excel_tabs_styles(self):
        path = self.make_excel(); result = self.run_tool('sheet-merge', files=[path, path])
        wb = load_workbook(result['files'][0]); self.assertEqual(len(wb.sheetnames), 2)
        self.assertTrue(wb.worksheets[1]['A1'].font.bold)
        self.assertEqual(wb.worksheets[1]['B2'].number_format, '$#,##0.00'); wb.close()

    def test_csv_formatted_excel_preserves_text_and_prevents_formula(self):
        source = self.root / 'Data.csv'; source.write_text('ID,Value\n00123,=1+1\n', encoding='utf-8')
        result = self.run_tool('sheet-excel', files=[source]); wb = load_workbook(result['files'][0]); ws = wb.active
        self.assertEqual(ws['A2'].value, '00123'); self.assertEqual(ws['B2'].data_type, 's')
        self.assertTrue(ws['A1'].font.bold); self.assertEqual(ws.freeze_panes, 'A2'); wb.close()

    def test_password_types(self):
        value = self.run_tool('password', files=[], kind='memorable')['password']
        self.assertEqual(len(value.split('-')), 6)
        value = self.run_tool('password', files=[], kind='pin', length=8)['password']
        self.assertTrue(value.isdigit()); self.assertEqual(len(value), 8)
        value = self.run_tool('password', files=[], kind='random', length=24)['password']
        self.assertEqual(len(value), 24); self.assertTrue(any(c.isupper() for c in value))

    def test_actual_mp4_to_mp3(self):
        import imageio_ffmpeg
        import subprocess
        ffmpeg = imageio_ffmpeg.get_ffmpeg_exe()
        video = self.root / 'Video.mp4'
        subprocess.run([ffmpeg, '-v', 'error', '-f', 'lavfi', '-i', 'color=c=blue:s=64x64:d=1', '-f', 'lavfi', '-i', 'sine=frequency=440:duration=1', '-c:v', 'mpeg4', '-c:a', 'aac', '-shortest', str(video)], check=True, capture_output=True)
        audio = self.run_tool('audio-mp3', files=[video])['files'][0]
        self.assertGreater(Path(audio).stat().st_size, 1000)
        subprocess.run([ffmpeg, '-v', 'error', '-i', audio, '-f', 'null', '-'], check=True, capture_output=True)

    def test_cancellation_cleans_partial_outputs(self):
        flag = self.root / 'cancel'; flag.touch()
        with self.assertRaises(worker.UserError):
            worker.dispatch({'tool':'pdf-split','files':[str(self.pdf)],'output_dir':str(self.out),'cancel_file':str(flag)})
        self.assertEqual(list(self.out.iterdir()), [])

    def test_speech_skips_unspoken_chunks_and_writes_audio(self):
        import numpy as np
        from types import SimpleNamespace
        chunks = []
        def create(phonemes, **options):
            self.assertTrue(phonemes)
            self.assertTrue(options['is_phonemes'])
            chunks.append(phonemes)
            return np.ones(2400, dtype=np.float32) * .1, 24000
        voice = SimpleNamespace(tokenizer=SimpleNamespace(phonemize=lambda text, **_: '' if text in ['---', '•', '\u200b'] else text), create=create)
        with patch('kokoro_onnx.Kokoro.from_session', return_value=voice), patch('onnxruntime.InferenceSession'), patch.object(worker, 'download', side_effect=lambda url, target, message: target), patch('imageio_ffmpeg.get_ffmpeg_exe', side_effect=RuntimeError('FFmpeg is not needed for speech')):
            result = self.run_tool('audio-speak', files=[], text='Hello.\n---\n•\n\u200b\nGoodbye.')
            with wave.open(result['files'][0]) as audio:
                self.assertEqual(audio.getnframes(), 4800)
                self.assertEqual(audio.getframerate(), 24000)
            self.assertEqual(chunks, ['Hello.', 'Goodbye.'])
            before = set(self.out.iterdir())
            with self.assertRaisesRegex(worker.UserError, 'no words'):
                self.run_tool('audio-speak', files=[], text='---\n•')
            self.assertEqual(set(self.out.iterdir()), before)

    def test_speech_error_is_specific_and_diagnostics_exclude_input(self):
        req = {'tool': 'audio-speak', 'options': {'text': 'PRIVATE TEXT'}, 'output_dir': str(self.out), 'model_dir': str(self.root / 'models')}
        with patch.object(worker, 'speech_task', side_effect=RuntimeError('PRIVATE TEXT')), patch('sys.stdin', io.StringIO(json.dumps(req))):
            worker.main()
        event = json.loads(worker.PROTOCOL.getvalue().splitlines()[-1])
        self.assertEqual(event['code'], 'speech-start-RuntimeError')
        self.assertNotIn('Check that it opens', event['message'])
        saved = (self.root / 'diagnostics/last-speech-error.json').read_text()
        self.assertNotIn('PRIVATE TEXT', saved)
        self.assertNotIn(str(self.root), saved)
        self.assertTrue(json.loads(saved)['frames'])

    def test_worker_reads_utf8_even_with_legacy_windows_stdio(self):
        text = '“Hello.” Café, naïve, ₹500, नमस्ते, 中文, 😊 and 𐐀.'
        request = {'tool': 'audio-speak', 'options': {'text': text}, 'files': [str(self.root / '“Résumé” 文書.pdf')]}
        payload = (json.dumps(request, ensure_ascii=False) + '\n').encode('utf-8')
        for encoding in ['cp1252', 'cp932', 'ascii']:
            with self.subTest(encoding=encoding):
                incoming = io.TextIOWrapper(io.BytesIO(payload), encoding=encoding, errors='surrogateescape')
                outgoing = io.TextIOWrapper(io.BytesIO(), encoding=encoding)
                errors = io.TextIOWrapper(io.BytesIO(), encoding=encoding)
                def process(decoded):
                    self.assertEqual(decoded, request)  # Never strip or replace the user's characters.
                    self.assertEqual(incoming.encoding, 'utf-8')
                    self.assertEqual(outgoing.encoding, 'utf-8')
                    self.assertEqual(errors.encoding, 'utf-8')
                    return {'text': decoded['options']['text']}
                worker.PROTOCOL = io.StringIO()
                with patch('sys.stdin', incoming), patch('sys.stdout', outgoing), patch('sys.stderr', errors), patch.object(worker, 'dispatch', side_effect=process):
                    worker.main()
                result = json.loads(worker.PROTOCOL.getvalue())
                self.assertEqual(result['event'], 'result')
                self.assertEqual(result['text'], text)
                incoming.close(); outgoing.close(); errors.close()


if __name__ == '__main__': unittest.main()
