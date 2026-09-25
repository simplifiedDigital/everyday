"""Everyday's local worker. One JSON request on stdin; JSON events on stdout."""
from __future__ import annotations

import base64
import contextlib
import copy
import csv
import hashlib
import io
import json
import math
import os
from pathlib import Path
import re
import secrets
import shutil
import stat
import string
import subprocess
import sys
import tempfile
import time
import urllib.request
from urllib.parse import quote
import wave
import zipfile

PROTOCOL = sys.stdout
REQUEST = {}
SPEECH_STAGE = "start"
MODEL_HASHES = {
    'kokoro-v1.0.int8.onnx': '6e742170d309016e5891a994e1ce1559c702a2ccd0075e67ef7157974f6406cb',
    'voices-v1.0.bin': 'bca610b8308e8d99f32e6fe4197e7ec01679264efed0cac9140fe9c29f1fbf7d',
    'config.json': '56a6d8110d311f19c8f0471e562832c7527f146b567275bfca59fcf7c184da9a',
    'model.bin': 'd01c3014881c9c6f3133c182f3d2887eb6ca1c789a7538c5c007196857a0a6a9',
    'tokenizer.json': 'fb7b63191e9bb045082c79fd742a3106a12c99513ab30df4a0d47fa6cb6fd0ab',
    'vocabulary.txt': '34ce3fe1c5041027b3f8d42912270993f986dbc4bb34cf27f951e34a1e453913',
}


class UserError(Exception):
    def __init__(self, message, code="invalid"):
        super().__init__(message)
        self.code = code


def emit(event, **data):
    PROTOCOL.write(json.dumps({"event": event, **data}, ensure_ascii=True) + "\n")
    PROTOCOL.flush()


def check_cancel():
    flag = REQUEST.get("cancel_file")
    if flag and Path(flag).exists():
        raise UserError("Cancelled. Your original files are unchanged.", "cancelled")


def progress(message, percent=None):
    check_cancel()
    emit("progress", message=message, percent=percent)


def safe_name(name):
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "-", name).strip(" .")[:120] or "result"


def output_folder(req):
    root = Path(req.get("output_dir") or Path.home() / "Downloads" / "Everyday").resolve()
    root.mkdir(parents=True, exist_ok=True)
    title = safe_name(Path(req.get("files", ["Everyday"])[0]).stem if req.get("files") else "Speech")
    folder = Path(tempfile.mkdtemp(prefix=title + "-", dir=root))
    return folder


def page_numbers(text, count):
    if not str(text).strip():
        return list(range(count))
    result = []
    try:
        for part in str(text).split(","):
            bounds = part.strip().split("-")
            if len(bounds) > 2:
                raise ValueError()
            a = int(bounds[0]); b = int(bounds[-1])
            if a < 1 or b > count or a > b:
                raise ValueError()
            result.extend(range(a - 1, b))
    except ValueError:
        raise UserError(f"Enter pages between 1 and {count}, for example 1-3, 5.")
    return list(dict.fromkeys(result))


def pdf_open(path, password=""):
    import pymupdf
    doc = pymupdf.open(path)
    if doc.needs_pass and not doc.authenticate(password):
        doc.close()
        raise UserError("This PDF needs its password.", "password")
    return doc


def inspect_file(req):
    from PIL import Image, ImageOps
    path = req["files"][0]
    if Path(path).suffix.lower() == ".pdf":
        import pymupdf
        with pdf_open(path, req.get("options", {}).get("password", "")) as doc:
            index = max(0, min(int(req.get("options", {}).get("page", 0)), len(doc) - 1))
            page = doc[index]
            pix = page.get_pixmap(matrix=pymupdf.Matrix(1.3, 1.3), alpha=False)
            return {"pages": len(doc), "page": index, "width": page.rect.width, "height": page.rect.height,
                    "preview": "data:image/png;base64," + base64.b64encode(pix.tobytes("png")).decode()}
    try:
        import pillow_heif
        pillow_heif.register_heif_opener()
        if req.get('options', {}).get('collage'):
            cols = max(1, min(4, int(req['options'].get('columns', 2))))
            count = min(40, len(req['files'])); cell = 160; gap = 4
            canvas = Image.new('RGB', (cols * cell + (cols + 1) * gap, math.ceil(count / cols) * cell + (math.ceil(count / cols) + 1) * gap), 'white')
            for i, item in enumerate(req['files'][:40]):
                check_cancel()
                with Image.open(item) as source:
                    thumb = ImageOps.fit(ImageOps.exif_transpose(source).convert('RGB'), (cell, cell))
                    canvas.paste(thumb, (gap + (i % cols) * (cell + gap), gap + (i // cols) * (cell + gap)))
            buf = io.BytesIO(); canvas.save(buf, 'JPEG', quality=85)
            return {'width': canvas.width, 'height': canvas.height, 'preview': 'data:image/jpeg;base64,' + base64.b64encode(buf.getvalue()).decode()}
        with Image.open(path) as src:
            img = ImageOps.exif_transpose(src)
            width, height = img.size
            img.thumbnail((900, 600))
            buf = io.BytesIO(); img.convert("RGB").save(buf, "JPEG", quality=85)
            return {"width": width, "height": height, "preview": "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode()}
    except (OSError, ValueError):
        return {}


def pdf_task(req, folder):
    import pymupdf
    op = req["tool"]; opts = req.get("options", {}); paths = req["files"]
    password = opts.get("password", "")
    if op == "pdf-merge":
        if len(paths) < 2:
            raise UserError("Add at least two PDFs to combine.")
        with pymupdf.open() as merged:
            for i, path in enumerate(paths):
                progress("Combining your PDFs…", i / len(paths) * 90)
                with pdf_open(path, password) as source:
                    merged.insert_pdf(source)
            target = folder / "Combined.pdf"
            merged.save(target, garbage=4, deflate=True)
        return {"files": [str(target)]}
    with pdf_open(paths[0], password) as doc:
        selected = page_numbers(opts.get("pages", ""), len(doc))
        stem = safe_name(Path(paths[0]).stem)
        target = folder / f"{stem}.pdf"
        if op == "pdf-split":
            files = []
            for i, n in enumerate(selected):
                progress("Saving your pages…", i / len(selected) * 95)
                dest = folder / f"{stem}-page-{n+1}.pdf"
                with pymupdf.open() as out:
                    out.insert_pdf(doc, from_page=n, to_page=n); out.save(dest)
                files.append(str(dest))
            return {"files": files}
        if op == "pdf-extract":
            n = int(opts.get("page", 1)) - 1
            if n not in range(len(doc)):
                raise UserError(f"Choose a page between 1 and {len(doc)}.")
            with pymupdf.open() as out:
                out.insert_pdf(doc, from_page=n, to_page=n); out.save(target)
        elif op == "pdf-rotate":
            for n in selected:
                doc[n].set_rotation((doc[n].rotation + int(opts.get("angle", 90))) % 360)
            doc.save(target, garbage=4)
        elif op == "pdf-protect":
            new_password = opts.get("newPassword", "")
            if len(new_password) < 1:
                raise UserError("Enter a password for your PDF.")
            if len(new_password.encode('utf-8')) > 40:
                raise UserError("Use a PDF password of up to 40 bytes (40 simple letters or numbers).")
            doc.save(target, encryption=pymupdf.PDF_ENCRYPT_AES_256,
                     owner_pw=secrets.token_urlsafe(24), user_pw=new_password,
                     permissions=pymupdf.PDF_PERM_PRINT | pymupdf.PDF_PERM_COPY | pymupdf.PDF_PERM_MODIFY | pymupdf.PDF_PERM_ANNOTATE)
        elif op == "pdf-unlock":
            if not password:
                raise UserError("Enter the current PDF password.", "password")
            doc.save(target, encryption=pymupdf.PDF_ENCRYPT_NONE, garbage=4)
        elif op == "pdf-redact":
            boxes = opts.get("boxes", [])
            if not boxes:
                raise UserError("Draw over the information you want to remove.")
            touched = set()
            for box in boxes:
                page = doc[int(box["page"])]
                r = page.rect
                rect = pymupdf.Rect(box["x"] * r.width, box["y"] * r.height,
                                    (box["x"] + box["w"]) * r.width, (box["y"] + box["h"]) * r.height)
                page.add_redact_annot(rect * page.derotation_matrix, fill=(0, 0, 0))
                touched.add(page.number)
            for n in touched:
                doc[n].apply_redactions(images=2, graphics=2, text=0)
            doc.scrub()
            doc.set_metadata({})
            doc.save(target, garbage=4, deflate=True, clean=True)
        elif op == "pdf-markdown":
            return markdown_task(doc, folder, stem)
        else:
            raise UserError("This PDF task is not available.")
        return {"files": [str(target)]}


def markdown_task(doc, folder, stem):
    import pymupdf4llm
    images = folder / "images"; images.mkdir()
    sections = []; scanned = 0
    for n, page in enumerate(doc):
        progress("Converting your document…", n / len(doc) * 95)
        # Keep a page image for scanned pages so no visual content silently disappears.
        if not page.get_text().strip():
            scanned += 1
            name = f"page-{n+1}.png"
            page.get_pixmap(dpi=150).save(images / name)
            sections.append(f"![Page {n+1}](images/{name})")
            continue
        with contextlib.redirect_stdout(sys.stderr):
            md = pymupdf4llm.to_markdown(doc, pages=[n], write_images=True,
                image_path=str(images), image_format="png", force_text=True, show_progress=False)
        # Converter emits absolute paths; use portable links beside the Markdown file.
        md = md.replace(str(images).replace("\\", "/") + "/", "images/").replace(str(images) + "/", "images/").replace(str(images) + "\\", "images/")
        md = re.sub(r'(!\[[^\]]*\]\()(images/[^)\n]+)(\))', lambda m: m[1] + quote(m[2], safe='/') + m[3], md)
        sections.append(md)
    target = folder / f"{stem}.md"
    target.write_text("\n\n".join(sections), encoding="utf-8")
    warning = f"{scanned} scanned page(s) were kept as images. Their text is not editable." if scanned else ""
    return {"files": [str(target)], "folder": str(folder), "warning": warning,
            "text": target.read_text(encoding="utf-8")[:100000]}


def image_task(req, folder):
    from PIL import Image, ImageOps
    import pillow_heif
    pillow_heif.register_heif_opener()
    opts = req.get("options", {}); op = req["tool"]
    fmt = opts.get("format", "jpg").lower()
    if fmt not in {"jpg", "png", "webp"}:
        raise UserError("Choose JPG, PNG, or WebP.")
    files = []; images = []
    if len(req['files']) > 40:
        raise UserError('Choose up to 40 photos at a time.')
    try:
        for i, path in enumerate(req["files"]):
            check_cancel()
            with Image.open(path) as source:
                images.append(ImageOps.exif_transpose(source).copy())
        if op == "image-collage":
            cols = max(1, min(4, int(opts.get("columns", 2))))
            gap = 16; cell = 640; rows = math.ceil(len(images) / cols)
            if rows > 20:
                raise UserError("Choose up to 40 photos for one collage.")
            canvas = Image.new("RGB", (cols * cell + (cols + 1) * gap, rows * cell + (rows + 1) * gap), "white")
            for i, img in enumerate(images):
                canvas.paste(ImageOps.fit(img.convert("RGB"), (cell, cell)), (gap + (i % cols) * (cell + gap), gap + (i // cols) * (cell + gap)))
            target = folder / "Collage.jpg"; canvas.save(target, quality=93); canvas.close()
            return {"files": [str(target)]}
        for i, img in enumerate(images):
            progress("Saving your photos…", i / len(images) * 95)
            if op == "image-resize":
                width = int(opts.get("width", 1200))
                if width < 1 or width > 16000:
                    raise UserError("Choose a width between 1 and 16,000 pixels.")
                img = img.resize((width, max(1, round(img.height * width / img.width))), Image.Resampling.LANCZOS)
            if fmt == "jpg":
                bg = Image.new("RGB", img.size, "white")
                if img.mode in ("RGBA", "LA") or "transparency" in img.info:
                    rgba = img.convert("RGBA"); bg.paste(rgba, mask=rgba.getchannel("A"))
                else:
                    bg.paste(img.convert("RGB"))
                img = bg
            target = folder / f"{safe_name(Path(req['files'][i]).stem)}-{i+1}.{fmt}"
            img.save(target, quality=92)
            files.append(str(target))
        return {"files": files}
    finally:
        for img in images:
            img.close()


def hash_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        while chunk := f.read(1024 * 1024):
            check_cancel(); h.update(chunk)
    return h.hexdigest()


def file_task(req, folder):
    paths = [Path(p) for p in req["files"]]; op = req["tool"]; opts = req.get("options", {})
    if op == "file-zip":
        target = folder / "Archive.zip"
        with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as archive:
            used = set()
            for i, path in enumerate(paths):
                progress("Creating your ZIP…", i / len(paths) * 95)
                name = path.name
                if name.lower() in used:
                    name = f"{i+1}-{name}"
                used.add(name.lower())
                archive.write(path, name)
        return {"files": [str(target)]}
    if op == "file-unzip":
        with zipfile.ZipFile(paths[0]) as archive:
            total = sum(x.file_size for x in archive.infolist())
            if total > 20 * 1024**3 or total > shutil.disk_usage(folder).free * .8:
                raise UserError("This ZIP needs more free space to unpack safely.")
            names = set()
            for info in archive.infolist():
                name = info.filename.replace("\\", "/")
                dest = (folder / name).resolve()
                if not dest.is_relative_to(folder.resolve()) or ":" in name or stat.S_ISLNK(info.external_attr >> 16):
                    raise UserError("This ZIP contains an unsafe file path.")
                if str(dest).lower() in names:
                    raise UserError("This ZIP contains conflicting file names.")
                names.add(str(dest).lower())
            for i, info in enumerate(archive.infolist()):
                progress("Unpacking your files…", i / max(1, len(names)) * 95)
                dest = folder / info.filename.replace("\\", "/")
                if info.is_dir():
                    dest.mkdir(parents=True, exist_ok=True)
                else:
                    dest.parent.mkdir(parents=True, exist_ok=True)
                    with archive.open(info) as source, open(dest, "xb") as out:
                        while chunk := source.read(1024 * 1024):
                            check_cancel(); out.write(chunk)
        return {"files": [str(p) for p in folder.rglob("*") if p.is_file()], "folder": str(folder)}
    if op == "file-split":
        size = int(opts.get("size", 100)) * 1024 * 1024
        if size < 1024 * 1024 or size > 4 * 1024**3:
            raise UserError("Choose a part size between 1 and 4096 MB.")
        source = paths[0]; parts = []; done = 0; overall = hashlib.sha256()
        with source.open("rb") as inp:
            while True:
                first = inp.read(min(size, 1024 * 1024))
                if not first: break
                part = folder / f"{safe_name(source.name)}.{len(parts)+1:03d}"
                h = hashlib.sha256(); written = 0
                with part.open("xb") as out:
                    chunk = first
                    while chunk:
                        check_cancel(); out.write(chunk); h.update(chunk); overall.update(chunk)
                        written += len(chunk); done += len(chunk)
                        progress("Splitting your file…", done / max(1, source.stat().st_size) * 95)
                        if written >= size: break
                        chunk = inp.read(min(1024 * 1024, size - written))
                parts.append({"name": part.name, "sha256": h.hexdigest()})
        manifest = folder / f"{safe_name(source.name)}.parts.json"
        manifest.write_text(json.dumps({"name": source.name, "sha256": overall.hexdigest(), "parts": parts}), encoding="utf-8")
        return {"files": [str(manifest)] + [str(folder / p["name"]) for p in parts], "folder": str(folder)}
    if op == "file-join":
        manifest_path = paths[0]; manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        target = folder / safe_name(manifest["name"])
        with target.open("xb") as out:
            for i, part in enumerate(manifest["parts"]):
                source = (manifest_path.parent / part["name"]).resolve()
                if source.parent != manifest_path.parent.resolve():
                    raise UserError("This parts list contains an invalid file name.")
                if not source.is_file() or hash_file(source) != part["sha256"]:
                    raise UserError(f"A part is missing or has changed: {part['name']}")
                progress("Putting your file back together…", i / max(1, len(manifest["parts"])) * 95)
                with source.open("rb") as inp:
                    while chunk := inp.read(1024 * 1024):
                        check_cancel(); out.write(chunk)
        if hash_file(target) != manifest["sha256"]:
            raise UserError("The joined file could not be verified.")
        return {"files": [str(target)]}
    raise UserError("Choose a file task.")


def spreadsheet_task(req, folder):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    op = req["tool"]; opts = req.get("options", {}); files = []
    if op == "sheet-csv":
        for path in req["files"]:
            wb = load_workbook(path, read_only=True, data_only=True)
            try:
                for sheet in wb:
                    progress(f"Saving {sheet.title}…")
                    target = folder / f"{safe_name(Path(path).stem)}-{safe_name(sheet.title)}-{len(files)+1}.csv"
                    with target.open("w", newline="", encoding="utf-8-sig") as out:
                        writer = csv.writer(out)
                        for row in sheet.values:
                            check_cancel(); writer.writerow(row)
                    files.append(str(target))
            finally:
                wb.close()
        return {"files": files, "warning": "Formulas use their last saved results. CSV does not include formatting."}
    wb = Workbook(); wb.remove(wb.active)
    if op == "sheet-merge":
        for path in req["files"]:
            progress(f"Adding {Path(path).name}…")
            src = load_workbook(path, data_only=True)
            try:
                for sheet in src:
                    title = re.sub(r'[\\*?:/\[\]]', "-", f"{Path(path).stem} - {sheet.title}")[:31]
                    out = wb.create_sheet(title)
                    for row in sheet:
                        check_cancel()
                        for cell in row:
                            new = out.cell(cell.row, cell.column, cell.value)
                            if cell.number_format: new.number_format = cell.number_format
                            # Style IDs belong to the source workbook; copy public components.
                            new.font = copy.copy(cell.font); new.fill = copy.copy(cell.fill)
                            new.border = copy.copy(cell.border); new.alignment = copy.copy(cell.alignment)
                            new.protection = copy.copy(cell.protection)
                    for key, dim in sheet.column_dimensions.items():
                        out.column_dimensions[key].width = dim.width
                    for key, dim in sheet.row_dimensions.items():
                        out.row_dimensions[key].height = dim.height
                    for merged in sheet.merged_cells.ranges: out.merge_cells(str(merged))
                    out.freeze_panes = sheet.freeze_panes
            finally:
                src.close()
        warning = "Combined saved values and basic formatting. Formulas, charts, and macros are not carried over."
    elif op == "sheet-excel":
        warning = ""
        for path in req["files"]:
            sheet = wb.create_sheet(re.sub(r'[\\*?:/\[\]]', "-", Path(path).stem)[:31])
            with open(path, encoding="utf-8-sig", newline="") as inp:
                sample = inp.read(8192); inp.seek(0)
                try: dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                except csv.Error: dialect = csv.excel
                for row in csv.reader(inp, dialect):
                    check_cancel()
                    sheet.append(row)
                    # CSV content is data, never executable spreadsheet formulas.
                    for cell in sheet[sheet.max_row]: cell.data_type = "s"
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="365F50")
                cell.alignment = Alignment(vertical="center")
            sheet.row_dimensions[1].height = 26
            sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
            for col in sheet.columns:
                width = min(48, max(12, max(len(str(c.value or "")) for c in col[:200]) + 2))
                sheet.column_dimensions[get_column_letter(col[0].column)].width = width
    else:
        raise UserError("Choose a spreadsheet task.")
    target = folder / ("Combined.xlsx" if op == "sheet-merge" else "Spreadsheet.xlsx")
    wb.save(target); wb.close()
    return {"files": [str(target)], "warning": warning}


def password_task(req):
    opts = req.get("options", {}); kind = opts.get("kind", "memorable")
    if kind == "memorable":
        words = (Path(__file__).parent / "words.txt").read_text(encoding="utf-8").split()
        count = max(4, min(10, int(opts.get("words", 6))))
        value = "-".join(secrets.choice(words) for _ in range(count))
    elif kind == "pin":
        value = "".join(secrets.choice(string.digits) for _ in range(max(6, min(32, int(opts.get("length", 8))))))
    else:
        size = max(8, min(128, int(opts.get("length", 20))))
        groups = [string.ascii_lowercase, string.ascii_uppercase, string.digits]
        if opts.get("symbols", True): groups.append("!@#$%&*+-=?")
        alphabet = "".join(groups)
        while True:
            value = "".join(secrets.choice(alphabet) for _ in range(size))
            if all(any(c in g for c in value) for g in groups): break
    return {"password": value}


def download(url, target, message):
    import ssl, certifi
    target = Path(target)
    expected = MODEL_HASHES.get(target.name)
    if target.exists():
        if expected is None or hash_file(target) == expected: return target
        target.unlink()
    target.parent.mkdir(parents=True, exist_ok=True)
    part = target.with_suffix(target.suffix + ".partial")
    progress(message)
    try:
        request = urllib.request.Request(url, headers={"User-Agent": "Everyday/0.1"})
        with urllib.request.urlopen(request, timeout=30, context=ssl.create_default_context(cafile=certifi.where())) as response, part.open("wb") as out:
            total = int(response.headers.get("Content-Length", 0)); count = 0; last = 0
            while chunk := response.read(1024 * 256):
                check_cancel(); out.write(chunk); count += len(chunk)
                if time.monotonic() - last > .25:
                    progress(message, count / total * 100 if total else None); last = time.monotonic()
        if expected and hash_file(part) != expected:
            raise UserError("The download could not be verified. Please try again.", "download")
        os.replace(part, target)
    except Exception:
        part.unlink(missing_ok=True)
        raise
    return target


def run_process(args):
    flags = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0
    with tempfile.TemporaryFile() as log:
        proc = subprocess.Popen(args, stdout=log, stderr=log, creationflags=flags)
        try:
            while proc.poll() is None:
                check_cancel(); time.sleep(.1)
            if proc.returncode:
                raise UserError("This audio file could not be read. Try another recording.")
        finally:
            if proc.poll() is None: proc.kill(); proc.wait()


def speech_task(req, folder):
    global SPEECH_STAGE
    models = Path(req.get("model_dir") or Path.home() / ".everyday" / "models")
    op = req["tool"]; opts = req.get("options", {})
    if op == "audio-mp3":
        import imageio_ffmpeg
        ffmpeg = imageio_ffmpeg.get_ffmpeg_exe()
        target = folder / f"{safe_name(Path(req['files'][0]).stem)}.mp3"
        progress("Creating your MP3…")
        run_process([ffmpeg, "-nostdin", "-y", "-i", req["files"][0], "-vn", "-codec:a", "libmp3lame", "-q:a", "2", str(target)])
        return {"files": [str(target)]}
    if op == "audio-transcribe":
        # CPU int8 keeps the model small and works without GPU drivers on both platforms.
        from faster_whisper import WhisperModel
        model_dir = models / "whisper-base"
        for name in ["config.json", "model.bin", "tokenizer.json", "vocabulary.txt"]:
            download("https://huggingface.co/Systran/faster-whisper-base/resolve/ebe41f70d5b6dfa9166e2c581c45c9c0cfc57b66/" + name,
                     model_dir / name, "Downloading speech tools. This only happens once.")
        progress("Listening to your recording…")
        model = WhisperModel(str(model_dir), device="cpu", compute_type="int8", cpu_threads=min(4, os.cpu_count() or 2), local_files_only=True)
        segments, info = model.transcribe(req["files"][0], beam_size=5, vad_filter=True)
        lines = []
        for segment in segments:
            check_cancel(); lines.append(segment.text.strip())
            progress("Writing your transcript…", min(99, segment.end / max(1, info.duration) * 100))
        value = "\n".join(lines)
        if not value.strip(): raise UserError("No speech was found in this recording.")
        target = folder / f"{safe_name(Path(req['files'][0]).stem)}.txt"
        target.write_text(value, encoding="utf-8")
        return {"files": [str(target)], "text": value}
    if op == "audio-speak":
        SPEECH_STAGE = "load"
        from kokoro_onnx import Kokoro
        import numpy as np
        text = opts.get("text", "").strip()
        if not text: raise UserError("Type or paste something to read aloud.")
        if len(text) > 50000: raise UserError("Use up to 50,000 characters at a time.")
        SPEECH_STAGE = "download"
        base = "https://github.com/thewh1teagle/kokoro-onnx/releases/download/model-files-v1.0/"
        name = "kokoro-v1.0.int8.onnx"
        model = download(base + name, models / name, "Downloading your voice. This only happens once.")
        voices = download(base + "voices-v1.0.bin", models / "voices-v1.0.bin", "Getting your voice ready…")
        progress("Reading your text…")
        SPEECH_STAGE = "voice"
        import onnxruntime as ort
        session_options = ort.SessionOptions()
        session_options.intra_op_num_threads = min(4, os.cpu_count() or 2)
        session_options.inter_op_num_threads = 1
        voice = Kokoro.from_session(ort.InferenceSession(str(model), sess_options=session_options, providers=['CPUExecutionProvider']), str(voices))
        target = folder / "Speech.wav"
        # Short chunks keep cancellation responsive and avoid model token truncation.
        import textwrap
        chunks = []
        for sentence in re.split(r'(?<=[.!?])\s+|\n+', text):
            chunks.extend(textwrap.wrap(sentence, width=350, break_long_words=True, break_on_hyphens=False))
        with wave.open(str(target), "wb") as wav:
            wav.setnchannels(1); wav.setsampwidth(2); wav.setframerate(24000)
            wrote_audio = False
            for i, chunk in enumerate(chunks):
                check_cancel()
                SPEECH_STAGE = "text"
                phonemes = voice.tokenizer.phonemize(chunk, lang="en-us")
                # Separators, bullets and invisible characters can produce no
                # phonemes. Kokoro otherwise crashes on np.concatenate([]).
                if phonemes.strip():
                    SPEECH_STAGE = "render"
                    samples, rate = voice.create(phonemes, voice="af_sarah", speed=1.0, lang="en-us", is_phonemes=True)
                    SPEECH_STAGE = "save"
                    wav.writeframes((np.clip(samples, -1, 1) * 32767).astype('<i2').tobytes())
                    wrote_audio = wrote_audio or len(samples) > 0
                progress("Reading your text…", (i + 1) / len(chunks) * 100)
        if not wrote_audio:
            raise UserError("There are no words to read aloud. Add some text and try again.", "speech-empty")
        return {"files": [str(target)]}
    raise UserError("Choose an audio task.")


def dispatch(req):
    global REQUEST, SPEECH_STAGE
    REQUEST = req
    SPEECH_STAGE = "start"
    tool = req.get("tool", "")
    if tool == "password": return password_task(req)
    if tool != "audio-speak" and not req.get("files"):
        raise UserError("Choose a file first.")
    if tool == "inspect": return inspect_file(req)
    folder = output_folder(req)
    try:
        if tool.startswith("pdf-"): result = pdf_task(req, folder)
        elif tool.startswith("image-"): result = image_task(req, folder)
        elif tool.startswith("file-"): result = file_task(req, folder)
        elif tool.startswith("sheet-"): result = spreadsheet_task(req, folder)
        elif tool.startswith("audio-"): result = speech_task(req, folder)
        else: raise UserError("Choose a tool to get started.")
        check_cancel()
        result["folder"] = str(folder)
        return result
    except BaseException:
        # This directory was freshly created by this job and never contains originals.
        shutil.rmtree(folder, ignore_errors=True)
        raise


def speech_failure(error):
    """Retain useful local diagnostics without text, paths or exception messages."""
    import traceback
    frames = [{"file": Path(frame.filename).name, "line": frame.lineno, "function": frame.name}
              for frame in traceback.extract_tb(error.__traceback__)]
    code = f"speech-{SPEECH_STAGE}-{type(error).__name__}"
    record = {"code": code, "stage": SPEECH_STAGE, "exception": type(error).__name__, "frames": frames}
    try:
        models = Path(REQUEST.get("model_dir") or Path.home() / ".everyday" / "models")
        diagnostics = models.parent / "diagnostics"
        diagnostics.mkdir(parents=True, exist_ok=True)
        (diagnostics / "last-speech-error.json").write_text(json.dumps(record, indent=2), encoding="utf-8")
    except OSError:
        pass  # A diagnostic write must never hide the original failure.
    return code


def main():
    try:
        # Both launchers send UTF-8 JSON. Frozen Python can ignore
        # PYTHONIOENCODING and otherwise decode stdin using the Windows locale.
        # That corrupts non-ASCII text (and can introduce surrogate characters).
        # Configure the streams inside the worker, before reading any input.
        for stream, errors in [(sys.stdin, "strict"), (sys.stdout, "backslashreplace"), (sys.stderr, "backslashreplace")]:
            if hasattr(stream, "reconfigure"):
                stream.reconfigure(encoding="utf-8", errors=errors)
        req = json.loads(sys.stdin.readline())
        with contextlib.redirect_stdout(sys.stderr):
            result = dispatch(req)
        emit("result", **result)
    except UserError as e:
        emit("error", message=str(e), code=e.code)
    except PermissionError:
        emit("error", message="This file or folder is in use. Close it and try again.", code="permission")
    except (urllib.error.URLError, TimeoutError, ConnectionError):
        emit("error", message="The download was interrupted. Check your connection and try again.", code="download")
    except Exception as e:
        if REQUEST.get("tool") == "audio-speak":
            code = speech_failure(e)
            emit("error", message=f"Couldn't create speech. Try again. If it happens again, share this code: {code}.", code=code)
            return
        print(type(e).__name__, str(e), file=sys.stderr)
        emit("error", message="We couldn't finish this file. Check that it opens normally, then try again.", code="failed")


if __name__ == "__main__":
    import multiprocessing
    multiprocessing.freeze_support()
    main()
